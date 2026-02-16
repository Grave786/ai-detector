from datetime import datetime
from io import BytesIO
import logging
import os
import pickle

import nltk
import numpy as np
from docx import Document
from fastapi import Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from pydantic import BaseModel
from pymongo import ReturnDocument
from nltk.tokenize import sent_tokenize
from pptx import Presentation

from backend.mongo import (
    ensure_collections_and_indexes,
    ensure_default_admin,
    scan_logs_collection,
    users_collection,
)
from backend.security import get_current_user
from features.feature_extractor import build_features_with_chunk_context
from router.admin import admin_router
from router.auth import auth_router

app = FastAPI(title="Turnitin-Style AI Detector")
logger = logging.getLogger("uvicorn.error")
CHUNK_SENTENCE_SIZE = int(os.getenv("CHUNK_SENTENCE_SIZE", "15"))

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://ai-checker-alpha.vercel.app",
        "http://127.0.0.1:5500",
        "http://localhost:5500",
        "http://127.0.0.1:5501",
        "http://localhost:5501",
        "http://127.0.0.1:3000",
        "http://localhost:3000",
        "http://127.0.0.1:5173",
        "http://localhost:5173",
    ],
    allow_origin_regex=r"https://.*\.vercel\.app|http://(localhost|127\.0\.0\.1|192\.168\.\d{1,3}\.\d{1,3}|10\.\d{1,3}\.\d{1,3}\.\d{1,3}|172\.(1[6-9]|2\d|3[0-1])\.\d{1,3}\.\d{1,3})(:\d+)?$",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.options("/{path:path}")
async def options_handler(path: str):
    return {"status": "ok"}


@app.get("/")
def health():
    return {"status": "running"}


app.include_router(auth_router)
app.include_router(admin_router)

ensure_collections_and_indexes()
ensure_default_admin()

try:
    nltk.data.find("tokenizers/punkt")
except LookupError:
    nltk.download("punkt", quiet=True)

MODEL_PATH = "models/xgb_model_.pkl"
if not os.path.exists(MODEL_PATH):
    raise FileNotFoundError("Model not found at models/xgb_model_.pkl")

with open(MODEL_PATH, "rb") as f:
    model = pickle.load(f)


class TextInput(BaseModel):
    text: str


def classify_turnitin(ai_percent, human_percent, polish_percent):
    if ai_percent >= 20 and ai_percent > human_percent:
        return "AI"

    if polish_percent >= 95:
        return "AI"

    if polish_percent < 90:
        return "AI" if ai_percent > human_percent else "Human"

    if human_percent >= ai_percent + 10:
        return "Human"

    if ai_percent >= 12:
        return "AI"

    return "Human"


def consume_user_token(current_user):
    user_doc = users_collection.find_one_and_update(
        {"_id": current_user["_id"], "tokens": {"$gt": 0}},
        {"$inc": {"tokens": -1}},
        return_document=ReturnDocument.BEFORE,
    )

    if not user_doc:
        existing = users_collection.find_one({"_id": current_user["_id"]})
        if not existing:
            raise HTTPException(status_code=403, detail="User Not Found")
        raise HTTPException(status_code=402, detail="TOKEN_FINISHED")

    return user_doc.get("tokens", 0)


def run_prediction(text: str, user_id: str, tokens_before: int):
    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail="Text is empty")

    sentences = sent_tokenize(text)
    if not sentences:
        raise HTTPException(status_code=400, detail="No sentences found")

    original_sentence_count = len(sentences)
    results = []
    total_ai = 0
    total_human = 0

    feature_rows = []
    sentence_order = []

    for i in range(0, len(sentences), CHUNK_SENTENCE_SIZE):
        chunk_sentences = sentences[i:i + CHUNK_SENTENCE_SIZE]
        chunk_text = " ".join(chunk_sentences)

        for sent in chunk_sentences:
            feature_rows.append(build_features_with_chunk_context(sent, chunk_text))
            sentence_order.append(sent)

    probs_batch = model.predict_proba(np.vstack(feature_rows))

    for sent, probs in zip(sentence_order, probs_batch):

        human_p = float(probs[0] * 100)
        ai_p = float(probs[1] * 100)
        polish_p = float(probs[2] * 100)

        total_ai += ai_p
        total_human += human_p

        final_label = classify_turnitin(ai_p, human_p, polish_p)
        results.append(
            {
                "sentence": sent,
                "human_probability": round(human_p, 2),
                "ai_probability": round(ai_p, 2),
                "over_polished_probability": round(polish_p, 2),
                "final_label": final_label,
            }
        )

    avg_ai = total_ai / len(sentences)
    avg_human = total_human / len(sentences)
    final_doc_label = "AI" if avg_ai > avg_human else "Human"

    scan_logs_collection.insert_one(
        {
            "uid": user_id,
            "scanned_text": text,
            "result": final_doc_label,
            "ai_percent": round(avg_ai, 2),
            "human_percent": round(avg_human, 2),
            "timestamp": datetime.utcnow(),
        }
    )

    return {
        "tokens_left": max(tokens_before - 1, 0),
        "overall_human_probability": round(avg_human, 2),
        "overall_ai_probability": round(avg_ai, 2),
        "final_document_label": final_doc_label,
        "sentences": results,
        "sentences_processed": len(sentences),
        "sentences_received": original_sentence_count,
    }


def extract_text_from_upload(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename or "")[1].lower()

    if ext == ".txt":
        try:
            return content.decode("utf-8")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="ignore")

    if ext == ".pdf":
        reader = PdfReader(BytesIO(content))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext in {".docx", ".word"}:
        document = Document(BytesIO(content))
        return "\n".join(p.text for p in document.paragraphs)

    if ext in {".ppt", ".pptx"}:
        try:
            presentation = Presentation(BytesIO(content))
        except Exception:
            raise HTTPException(
                status_code=400,
                detail="Unable to read .ppt file. Please convert it to .pptx and retry.",
            )

        lines = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    lines.append(shape.text)
        return "\n".join(lines)

    if ext == ".doc":
        raise HTTPException(
            status_code=400,
            detail="Legacy .doc is not supported directly. Please convert to .docx.",
        )

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Use txt, pdf, docx, pptx, or ppt.",
    )


@app.post("/predict")
def predict(data: TextInput, current_user=Depends(get_current_user)):
    user_id = str(current_user["_id"])
    text = data.text.strip()
    tokens_before = consume_user_token(current_user)
    return run_prediction(text=text, user_id=user_id, tokens_before=tokens_before)


@app.post("/predict-file")
async def predict_file(file: UploadFile = File(...), current_user=Depends(get_current_user)):
    user_id = str(current_user["_id"])

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    extracted_text = extract_text_from_upload(file.filename or "", content).strip()
    if not extracted_text:
        raise HTTPException(status_code=400, detail="No readable text found in file")

    tokens_before = consume_user_token(current_user)
    return run_prediction(text=extracted_text, user_id=user_id, tokens_before=tokens_before)


@app.post("/extract-file")
async def extract_file(file: UploadFile = File(...), _current_user=Depends(get_current_user)):
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")

    extracted_text = extract_text_from_upload(file.filename or "", content).strip()
    if not extracted_text:
        raise HTTPException(status_code=400, detail="No readable text found in file")

    return {
        "filename": file.filename or "",
        "text": extracted_text,
        "characters": len(extracted_text),
    }


@app.get("/my-history")
def get_my_history(current_user=Depends(get_current_user)):
    user_id = str(current_user["_id"])
    logs = scan_logs_collection.find({"uid": user_id}).sort("timestamp", -1).limit(100)

    return [
        {
            "id": str(log["_id"]),
            "scanned_text": log.get("scanned_text", ""),
            "result": log.get("result"),
            "ai_percent": log.get("ai_percent", 0),
            "human_percent": log.get("human_percent", 0),
            "timestamp": log.get("timestamp").isoformat() if log.get("timestamp") else None,
        }
        for log in logs
    ]


